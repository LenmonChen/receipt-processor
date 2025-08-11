# 把收到的图片格式的文件,转换成word模式, 为了满足Qwen-long模型的需求
from docx import Document
from docx.shared import Inches

def img_changed_to_word(image_path, new_doc_path):
    # 创建一个新的 Word 文档 或 打开已有的文档
    document = Document()  # 新建一个文档
    # 或者如果你想在已有文档中添加图片
    # document = Document('existing_document.docx')

    # 添加图片到文档中
    # 参数：图片路径、宽度（可选）
    # image_path = r'C:\Users\13916\Desktop\wechat_2025-07-30_190907_709.png'  # 替换为你的图片路径
    document.add_picture(image_path, width=Inches(4))  # 宽度为 4 英寸

    # 保存文档
    document.save(new_doc_path)
    print("图片已成功插入到 Word 文档中！")
